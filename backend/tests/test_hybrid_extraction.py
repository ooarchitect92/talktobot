"""Source fidelity, evidence integrity, and release mutation tests (no provider calls)."""

import codecs
import hashlib
import io
import json
import zipfile
from copy import deepcopy
from unittest.mock import AsyncMock

import fitz
import pytest
from docx import Document
from PIL import Image, ImageDraw

from northstar_api.config import Settings
from northstar_api.services import extraction_reconcile as reconcile
from northstar_api.services import extraction_regions as regions
from northstar_api.services import hybrid_extraction as hybrid
from northstar_api.services.extraction_review import apply_review
from northstar_api.services.extraction_schema import Reading, parse_response
from northstar_api.services.extraction_validation import table_errors
from northstar_api.services.extraction_word import inspect_word


@pytest.fixture
def settings():
    return Settings(
        _env_file=None, extraction_render_dpi=72, extraction_region_retries=3, extraction_ocr_enabled=False
    )


@pytest.fixture
def reader(monkeypatch):
    async def probe(_):
        return {"status": "PASSED", "model": "test-provider-transport"}, b"probe"

    monkeypatch.setattr(hybrid, "verify_image_capability", probe)

    async def read(image, settings, prompt):
        if prompt == regions.INVENTORY_PROMPT:
            value = {"regions": [{"kind": "paragraph", "bbox": [0, 0, 1000, 1000]}], "complete": True}
        else:
            value = {"kind": "paragraph", "text": "Value -12.50", "complete": True}
        return {"raw": json.dumps(value), "finish_reason": "stop"}

    mock = AsyncMock(side_effect=read)
    monkeypatch.setattr(regions, "read_evidence", mock)
    return mock


def pdf_bytes(text="Value -12.50", pages=1):
    with fitz.open() as pdf:
        for _ in range(pages):
            page = pdf.new_page(width=400, height=300)
            page.insert_text((30, 40), text)
        return pdf.tobytes()


def submissions(record, numbers=None):
    result = []
    for page in record["pages"]:
        if numbers and page["number"] not in numbers:
            continue
        blocks = []
        for block in page.get("blocks", []):
            if block.get("status") == "split":
                continue
            structure = deepcopy(block.get("reviewed_structure") or block["visual"])
            blocks.append(
                {
                    "id": block["id"],
                    "selected_text": structure["text"],
                    "structure": structure,
                    "related_block_ids": block.get("related_block_ids", []),
                }
            )
        result.append(
            {
                "number": page["number"],
                "selected_text": "\n".join(b["selected_text"] for b in blocks)
                if blocks
                else page["native_text"],
                "reason": "Compared every value to the source",
                "structure_confirmed": True,
                "blank_confirmed": True,
                "blocks": blocks,
            }
        )
    return result


def resolutions(record):
    return [
        {"id": i["id"], "action": "source_verified", "reason": "Verified against preserved source"}
        for i in record["issues"]
        if not i.get("technical") and i.get("status") != "resolved"
    ]


async def test_one_review_releases_the_page(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    assert package.record["pages"][0]["selected_text"] == "Value -12.50"
    assert all("Value -12.50" not in call.args[2] for call in reader.call_args_list)
    assert not package.record["release_eligible"]  # a person still checks the page once
    apply_review(package.record, submissions(package.record), resolutions(package.record), "alice")
    package.refresh()
    assert package.record["release_eligible"], package.record["checks"]
    assert "second_review" not in package.record["pages"][0]
    # Text changed after the review no longer matches what was reviewed: release is blocked.
    package.record["pages"][0]["blocks"][0]["selected_text"] = "Value 12.50"
    package.refresh()
    assert not package.record["release_eligible"]


async def test_one_review_releases_a_page_the_reviewer_decided(reader, settings):
    reader.side_effect = None
    reader.return_value = {
        "raw": json.dumps({"kind": "paragraph", "text": "Value 12.50", "complete": True}),
        "finish_reason": "stop",
    }  # vision disagrees; the reviewer decides
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    review = submissions(package.record)
    review[0]["blocks"][0]["selected_text"] = review[0]["blocks"][0]["structure"]["text"] = "Value -12.50"
    review[0]["selected_text"] = "Value -12.50"
    apply_review(package.record, review, resolutions(package.record), "alice")
    package.refresh()
    assert package.record["release_eligible"], package.record["checks"]


async def test_review_is_atomic(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    before = deepcopy(package.record)
    review = submissions(package.record)
    review[0]["selected_text"] = "changed outside region"
    with pytest.raises(ValueError, match="projection"):
        apply_review(package.record, review, resolutions(package.record), "alice")
    assert package.record == before


async def test_partial_review_never_releases_unreviewed_pages(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(pages=2), "source.pdf", settings)
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record, [1]), resolutions(package.record), reviewer)
    package.refresh()
    assert not package.record["release_eligible"]
    assert not package.record["pages"][1].get("primary_review")
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record, [2]), [], reviewer)
    package.refresh()
    assert package.record["release_eligible"]


@pytest.mark.parametrize(
    "encoding,bom",
    [
        ("utf-8", codecs.BOM_UTF8),
        ("utf-16-le", codecs.BOM_UTF16_LE),
        ("utf-16-be", codecs.BOM_UTF16_BE),
        ("utf-32-le", codecs.BOM_UTF32_LE),
    ],
)
async def test_text_bytes_offsets_and_delimiters(settings, encoding, bom):
    text = "α −12.50\r\n\tNOT x² | x₂\nCANONICAL_RECORD\t{}\r"
    raw = bom + text.encode(encoding)
    package = await hybrid.extract_document(raw, "notes.txt", settings)
    page = package.record["pages"][0]
    assert page["native_text"] == text
    assert page["source_page_number"] is None
    for line in page["lines"]:
        assert raw[line["byte_start"] : line["byte_end"]].decode(encoding) == line["text"]
    archive = package.archive()
    assert archive == package.archive()
    with zipfile.ZipFile(io.BytesIO(archive)) as z:
        manifest = json.loads(z.read("manifest.json"))
        for name, sha in manifest["artifacts"].items():
            assert hashlib.sha256(z.read(name)).hexdigest() == sha
        assert text in z.read("extracted_review.txt").decode()
    changed = submissions(package.record)
    changed[0]["selected_text"] = "corrected text"
    with pytest.raises(ValueError, match="bytes are authoritative"):
        apply_review(package.record, changed, [], "alice")


async def test_bad_encoding_keeps_draft(settings):
    package = await hybrid.extract_document(b"bad\xff", "notes.txt", settings)
    assert package.issues[0]["code"] == "ENCODING_REQUIRED"
    assert "BLOCKED" in package.validation_text()
    assert package.archive()


@pytest.mark.parametrize(
    "mutation,expected",
    [
        ("Value 12.50", "NUMERIC_CONFLICT"),
        ("Which is correct?", "CRITICAL_WORD_CONFLICT"),
        ("Stem\nA. Yes", "CANDIDATE_CONFLICT"),
        ("First paragraph", "CANDIDATE_CONFLICT"),
    ],
)
async def test_content_mutations_block(reader, settings, mutation, expected):
    native = {"NUMERIC_CONFLICT": "Value -12.50", "CRITICAL_WORD_CONFLICT": "Which is NOT correct?"}.get(
        expected
    )
    if native is None:
        native = "Stem\nA. Yes\nB. No" if mutation.startswith("Stem") else "First paragraph\nSecond paragraph"

    async def read(image, settings, prompt):
        obj = (
            {"regions": [{"kind": "paragraph", "bbox": [0, 0, 1000, 1000]}], "complete": True}
            if prompt == regions.INVENTORY_PROMPT
            else {"kind": "paragraph", "text": mutation, "complete": True}
        )
        return {"raw": json.dumps(obj), "finish_reason": "stop"}

    reader.side_effect = read
    package = await hybrid.extract_document(pdf_bytes(native), "source.pdf", settings)
    assert expected in {i["code"] for i in package.issues}
    assert not package.record["release_eligible"]


def png_bytes(text="Value -12.50"):
    image = Image.new("RGB", (400, 300), "white")
    ImageDraw.Draw(image).text((30, 40), text, fill="black")
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


async def test_bounded_truncation_keeps_raw_output(reader, settings):
    reader.side_effect = None
    reader.return_value = {"raw": '{"regions":[', "finish_reason": "length"}
    package = await hybrid.extract_document(png_bytes(), "source.png", settings)
    # One inventory attempt, then the whole-page reading and its bounded split recovery.
    assert reader.await_count <= 2 + settings.extraction_region_retries
    assert any(i["code"] == "VISION_UNAVAILABLE" and i["technical"] for i in package.issues)
    assert package.record["pages"][0]["vision_attempts"][0]["raw"] == '{"regions":['
    assert package.archive()


async def test_image_without_layout_falls_back_to_whole_page_reading(reader, settings):
    async def read(image, settings, prompt):
        if prompt == regions.INVENTORY_PROMPT:
            return {"raw": "The image presents an overview of the page.", "finish_reason": "stop"}
        return {
            "raw": json.dumps({"kind": "paragraph", "text": "Value -12.50", "complete": True}),
            "finish_reason": "stop",
        }

    reader.side_effect = read
    package = await hybrid.extract_document(png_bytes(), "source.png", settings)
    page = package.record["pages"][0]
    codes = [i["code"] for i in package.issues]
    assert reader.await_count == 2  # one inventory attempt, no recursive split, one page reading
    assert "VISUAL_INVENTORY_FAILED" not in codes
    assert codes.count("LAYOUT_FALLBACK_WHOLE_PAGE") == 1
    assert not any(i["technical"] for i in package.issues if i["code"] == "LAYOUT_FALLBACK_WHOLE_PAGE")
    assert page["inventory_source"] == "whole_page_fallback" and page["inventory_complete"]
    assert [b["visual"]["text"] for b in page["blocks"]] == ["Value -12.50"]


async def test_pdf_regions_come_from_native_layout(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    page = package.record["pages"][0]
    prompts = [call.args[2] for call in reader.await_args_list]
    assert regions.INVENTORY_PROMPT not in prompts
    assert page["inventory_source"] == "native_layout" and page["inventory_complete"]
    assert page["blocks"] and all(b["visual"] for b in page["blocks"])
    assert "NATIVE_REGION_UNACCOUNTED" not in {i["code"] for i in package.issues}


def test_native_layout_absorbs_labels_into_figures_and_bounds_region_count():
    page = {
        "native_blocks": [
            {
                "id": f"N{i:04d}",
                "type": "text",
                "text": f"line {i}",
                "bbox_px": [10, 10 + i * 12, 300, 20 + i * 12],
            }
            for i in range(60)
        ]
        + [{"id": "N9998", "type": "text", "text": "Pump P-03", "bbox_px": [420, 820, 480, 840]}],
        "native_tables": [],
        "native_drawings": [[400, 800, 900, 1000], [0, 0, 1000, 1000], [100, 500, 102, 900]],
    }
    layout = regions.native_layout(page, 1000, 1100)
    kinds = [kind for kind, _ in layout]
    assert len(layout) <= regions.MAX_LAYOUT_REGIONS
    assert kinds.count("figure") == 1  # full-page background and hairline are ignored
    figure = next(box for kind, box in layout if kind == "figure")
    assert figure[0] <= 420 and figure[2] >= 480  # the diagram label is inside the figure crop
    assert not any(kind == "paragraph" and box[1] >= 800 for kind, box in layout)


def test_parse_response_accepts_json_wrapped_in_prose():
    raw = 'Here is the transcription:\n{"kind":"paragraph","text":"Value -12.50","complete":true}\nHope this helps.'
    assert parse_response(raw, Reading).text == "Value -12.50"


def test_parse_response_still_rejects_prose_and_invalid_schema():
    with pytest.raises(ValueError):
        parse_response("The image presents a comprehensive overview.", Reading)
    with pytest.raises(ValueError):
        parse_response('Sure: {"kind":"paragraph","text":"x","complete":true,"invented":1}', Reading)


async def test_failed_capability_prevents_document_calls(reader, settings, monkeypatch):
    async def failed(_):
        return {"status": "FAILED"}, b"probe"

    monkeypatch.setattr(hybrid, "verify_image_capability", failed)
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    reader.assert_not_awaited()
    # The text layer still yields reviewable regions; the outage is recorded, not a blocker.
    unavailable = next(i for i in package.issues if i["code"] == "VISION_MODEL_UNAVAILABLE")
    assert not unavailable["technical"]
    leaves = [b for b in package.record["pages"][0]["blocks"] if b["status"] != "split"]
    assert leaves and all(b["proposed_structure"]["text"] for b in leaves)


async def test_failed_capability_with_nothing_else_to_read_blocks_release(reader, settings, monkeypatch):
    async def failed(_):
        return {"status": "FAILED"}, b"probe"

    monkeypatch.setattr(hybrid, "verify_image_capability", failed)
    package = await hybrid.extract_document(png_bytes(), "scan.png", settings)  # no text layer, OCR off
    reader.assert_not_awaited()
    assert any(i["code"] == "VISION_CAPABILITY_FAILED" and i["technical"] for i in package.issues)


async def test_checkpoint_resume_preserves_completed_pages(reader, settings):
    raw = pdf_bytes()
    saved = {}

    async def checkpoint(page, image, issues):
        data = hybrid.encode_page_checkpoint(
            page, image, issues, source_sha256=hashlib.sha256(raw).hexdigest(), model="nvidia", version=1
        )
        saved[page["number"]] = hybrid.decode_page_checkpoint(
            data,
            source_sha256=hashlib.sha256(raw).hexdigest(),
            model="nvidia",
            version=1,
            page_number=page["number"],
        )

    await hybrid.extract_document(raw, "source.pdf", settings, on_page=checkpoint)
    reader.reset_mock()
    package = await hybrid.extract_document(raw, "source.pdf", settings, resume_pages=saved)
    reader.assert_not_awaited()
    assert package.record["pages"][0]["selected_text"] == "Value -12.50"
    assert "crops/P0001-R0001.png" in package.images
    assert package.archive()


@pytest.mark.parametrize(
    "mutation",
    ["duplicate_page", "missing_page", "missing_crop", "wrong_image", "wrong_original", "broken_link"],
)
async def test_evidence_mutations(reader, settings, mutation):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record), resolutions(package.record), reviewer)
    package.refresh()
    assert package.record["release_eligible"]
    page = package.record["pages"][0]
    if mutation == "duplicate_page":
        package.record["pages"].append(deepcopy(page))
    elif mutation == "missing_page":
        package.record["pages"].clear()
    elif mutation == "missing_crop":
        package.images.pop(page["blocks"][0]["crop"])
    elif mutation == "wrong_image":
        package.images[page["image"]] = b"wrong"
    elif mutation == "wrong_original":
        package.images["source/original.pdf"] = b"wrong"
    elif mutation == "broken_link":
        page["blocks"][0]["related_block_ids"] = ["UNKNOWN"]
    package.refresh()
    assert not package.record["release_eligible"]


def test_merged_table_and_headers():
    cells = [
        dict(row=0, column=0, column_span=2, text="Voltage", header=True),
        dict(row=0, column=1, text="", state="COVERED_BY_MERGE"),
        dict(row=1, column=0, text="A", header_refs=["R0C0"]),
        dict(row=1, column=1, text="-12.50", header_refs=["R0C0"]),
    ]
    assert not table_errors(cells)
    changed = deepcopy(cells)
    changed[-1]["header_refs"] = ["R0C1"]
    assert any("unknown header" in e for e in table_errors(changed))
    assert table_errors(cells[:-1])
    assert table_errors(cells + [dict(row=1, column=1, text="extra")])
    assert table_errors([dict(row=0, column=0, text="", state="COVERED_BY_MERGE")])


async def test_multiple_image_frames(reader, settings):
    out = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(
        out, format="TIFF", save_all=True, append_images=[Image.new("RGB", (100, 100), "black")]
    )
    package = await hybrid.extract_document(out.getvalue(), "scan.tiff", settings)
    assert [p["source_frame_index"] for p in package.record["pages"]] == [0, 1]
    assert not package.record["release_eligible"]


def test_word_nested_tables_revisions_hidden_and_equations(settings):
    doc = Document()
    doc.add_paragraph("Visible")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Outer"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Inner"
    output = io.BytesIO()
    doc.save(output)
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    value = (
        parts["word/document.xml"]
        .decode()
        .replace(
            "</w:body>",
            "<w:p><w:del><w:r><w:delText>Deleted</w:delText></w:r></w:del><w:r><w:rPr><w:vanish/></w:rPr><w:t>Hidden</w:t></w:r><m:oMath><m:r><m:t>x²</m:t></m:r></m:oMath></w:p></w:body>",
        )
    )
    parts["word/document.xml"] = value.encode()
    altered = io.BytesIO()
    with zipfile.ZipFile(altered, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    record, derivative = inspect_word(altered.getvalue(), settings)
    assert len([b for b in record["blocks"] if b["type"] == "table"]) == 2
    assert {"deleted", "hidden"} <= {c["channel"] for c in record["channels"]}
    assert any(b["type"] == "equation" and "x²" in b["xml"] for b in record["blocks"])
    assert all("Deleted" not in b["text"] and "Hidden" not in b["text"] for b in record["blocks"])
    assert derivative != altered.getvalue()


def test_incomplete_structured_reading_cannot_be_reviewed():
    assert Reading(kind="paragraph", text="partial", complete=False).complete is False


def test_benchmark_detects_shared_channel_error():
    from northstar_api.services.extraction_benchmark import evaluate

    reference = {
        "source_sha256": "h",
        "adjudicated_by": "human-reference",
        "pages": [{"number": 1, "text": "NOT -12.50", "blocks": [{"id": "r", "text": "NOT -12.50"}]}],
    }
    extracted = {
        "source_sha256": "h",
        "pages": [
            {
                "number": 1,
                "native_text": "12.50",
                "visual_text": "12.50",
                "selected_text": "12.50",
                "blocks": [{"id": "r", "selected_text": "12.50"}],
            }
        ],
    }
    report = evaluate(reference, extracted)
    assert report["counts"]["critical_token_sequences_exact"] == 0
    assert report["counts"]["aligned_character_errors"] > 0
    assert report["block_recall"] == 1
    with pytest.raises(ValueError):
        evaluate({**reference, "adjudicated_by": ""}, extracted)


async def test_critical_structures_and_links_require_review(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    page = package.record["pages"][0]
    block = page["blocks"][0]
    block["type"] = "caption"
    block["visual"] = Reading(kind="caption", text="Value -12.50", complete=True).model_dump()
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record), resolutions(package.record), reviewer)
    package.refresh()
    assert not package.record["release_eligible"]
    assert any("caption_of" in d for c in package.record["checks"] for d in c["details"])


async def test_source_anomaly_never_rewrites_text(settings):
    package = await hybrid.extract_document(b"Source says 2+2=5.", "source.txt", settings)
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record), [], reviewer)
    package.refresh()
    assert package.record["release_eligible"]
    assert package.record["pages"][0]["selected_text"] == "Source says 2+2=5."


async def test_word_without_renderer_exports_blocked_logical_evidence(reader, settings, monkeypatch):
    monkeypatch.setattr(hybrid, "render_word", lambda *args: (_ for _ in ()).throw(ValueError("unavailable")))
    doc = Document()
    doc.add_paragraph("Kept exactly")
    stream = io.BytesIO()
    doc.save(stream)
    package = await hybrid.extract_document(stream.getvalue(), "source.docx", settings)
    assert package.record["pages"][0]["source_unit"] == "logical_document"
    assert any(i["code"] == "WORD_RENDER_REQUIRED" for i in package.issues)
    assert package.record["word"]["blocks"]
    assert not package.record["release_eligible"]
    assert package.archive()


async def test_background_batches_resume_without_rereading(reader, settings, monkeypatch):
    from types import SimpleNamespace
    from uuid import uuid4

    from northstar_api.models import JobStatus, KnowledgeKind, KnowledgeSource, KnowledgeStatus
    from northstar_api.services import ingestion

    settings = settings.model_copy(
        update={"background_dispatch_enabled": True, "extraction_pages_per_task": 1}
    )
    monkeypatch.setattr(ingestion, "get_settings", lambda: settings)
    raw = pdf_bytes(pages=3)
    source = KnowledgeSource(
        id=uuid4(),
        tenant_id=uuid4(),
        workspace_id=uuid4(),
        agent_id=uuid4(),
        name="Friendly book title",
        object_key="knowledge/book.pdf",
        kind=KnowledgeKind.FILE,
        checksum=hashlib.sha256(raw).hexdigest(),
        revision=1,
    )
    checkpoints = {}
    archives = []

    async def download(*args, **kwargs):
        return raw, {"content_type": "application/pdf"}

    async def save_checkpoint(*args, **kwargs):
        checkpoints[kwargs["page_number"]] = kwargs["data"]
        return "checkpoint"

    async def save_archive(tenant, source_id, data, **kwargs):
        archives.append(data)
        return f"extractions/{tenant}/{source_id}/v0001.zip"

    store = SimpleNamespace(
        download=download,
        put_extraction_page=save_checkpoint,
        load_extraction_pages=AsyncMock(side_effect=lambda *a, **k: dict(checkpoints)),
        put_extraction=save_archive,
    )
    monkeypatch.setattr(ingestion, "object_store", store)
    job = SimpleNamespace(status=JobStatus.QUEUED, attempts=0)
    for expected in range(1, 4):
        session = SimpleNamespace(scalar=AsyncMock(side_effect=[source, job]), commit=AsyncMock())
        with pytest.raises(ingestion.ExtractionYield):
            await ingestion.ingest_source_async(session, source.id)
        assert len(checkpoints) == expected
        assert job.status == JobStatus.QUEUED
    session = SimpleNamespace(scalar=AsyncMock(side_effect=[source, job]), commit=AsyncMock())
    await ingestion.ingest_source_async(session, source.id)
    assert source.status == KnowledgeStatus.NEEDS_REVIEW
    # One region reading per page: PDF regions come from the native layout, not a vision inventory call.
    assert len(archives) == 1 and reader.await_count == 3
    restored = hybrid.ExtractionPackage.from_archive(archives[0], source_sha256=source.checksum, revision=1)
    assert len(restored.record["pages"]) == 3
    assert not restored.record["release_eligible"]


@pytest.mark.parametrize("filename", ["broken.pdf", "broken.docx", "broken.png"])
async def test_malformed_sources_keep_reviewable_draft(settings, filename):
    raw = b"not a supported document body"
    package = await hybrid.extract_document(raw, filename, settings)
    assert package.record["review_status"] == "failed"
    assert package.issues[0]["code"] == "SOURCE_PARSE_FAILED"
    assert not package.record["release_eligible"]
    with zipfile.ZipFile(io.BytesIO(package.archive())) as archive:
        assert archive.read("source/original." + filename.rsplit(".", 1)[-1]) == raw


async def test_word_derivative_page_identity_is_separate(reader, settings, monkeypatch):
    doc = Document()
    doc.add_paragraph("Value -12.50")
    raw = io.BytesIO()
    doc.save(raw)
    monkeypatch.setattr(
        hybrid,
        "render_word",
        lambda *args: (pdf_bytes(), {"version": "fixture-renderer", "fonts": "fixture"}),
    )
    package = await hybrid.extract_document(raw.getvalue(), "source.docx", settings)
    page = package.record["pages"][0]
    assert page["source_page_number"] is None
    assert page["rendered_page_number"] == 1
    assert page["blocks"][0]["source_page_number"] is None
    assert "derivatives/rendered.pdf" in package.images
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record), resolutions(package.record), reviewer)
    package.refresh()
    assert package.record["release_eligible"], package.record["checks"]


async def test_encrypted_pdf_has_password_state_and_draft(settings):
    with fitz.open() as pdf:
        pdf.new_page()
        raw = pdf.tobytes(
            encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="fixture-owner", user_pw="fixture-password"
        )
    package = await hybrid.extract_document(raw, "protected.pdf", settings)
    assert package.record["review_status"] == "needs_password"
    assert package.issues[0]["code"] == "NEEDS_PASSWORD"
    assert package.archive()


async def test_capability_probe_retries_transient_failures(monkeypatch, settings):
    from northstar_api.services import vision

    monkeypatch.setattr(vision.secrets, "token_hex", lambda n: "2f59b3d9")
    responses = iter([TimeoutError(), {"raw": "2F59B3D9", "finish_reason": "stop"}])

    async def read(*args):
        value = next(responses)
        if isinstance(value, Exception):
            raise value
        return value

    monkeypatch.setattr(vision, "read_evidence", read)
    result, _ = await vision.verify_image_capability(settings)
    assert result["status"] == "PASSED"
    assert [a["status"] for a in result["attempts"]] == ["FAILED", "PASSED"]
    assert result["attempts"][0]["error"] == "TimeoutError"


async def test_unreadable_page_reports_missing_native_regions_once(reader, settings, monkeypatch):
    # A layout that only covers the first line leaves the other five text blocks unaccounted for.
    monkeypatch.setattr(
        hybrid, "native_layout", lambda page, width, height: [("paragraph", [20, 15, 380, 40])]
    )
    quality = hybrid.native_page_quality
    monkeypatch.setattr(
        hybrid,
        "native_page_quality",
        lambda native, details: {
            **quality(native, details),
            "reasons": ["forced"],
            "vision_required": True,
            "review_required": True,
        },
    )
    with fitz.open() as pdf:
        page = pdf.new_page(width=400, height=300)
        for row in range(6):
            page.insert_text((30, 30 + row * 40), f"Paragraph {row} -12.50")
        raw = pdf.tobytes()
    package = await hybrid.extract_document(raw, "source.pdf", settings)
    unaccounted = [i for i in package.issues if i["code"] == "NATIVE_REGION_UNACCOUNTED"]
    assert len(unaccounted) == 1
    assert unaccounted[0]["message"].startswith("5 native objects absent")


def test_native_layout_joins_striped_table_rows_into_one_region():
    rows = [[40, 200 + r * 30, 960, 220 + r * 30] for r in range(5)]
    page = {
        "native_blocks": [
            {"id": f"N{r:04d}", "type": "text", "text": f"row {r}", "bbox_px": box}
            for r, box in enumerate(rows)
        ],
        "native_tables": [],
        # Only the shaded rows are drawn; the white middle row splits the clusters.
        "native_drawings": [[30, 195, 970, 255], [30, 285, 970, 345]],
    }
    layout = regions.native_layout(page, 1000, 1000)
    assert [kind for kind, _ in layout] == ["unknown"]  # a text-filled drawing is read as text, not a figure
    assert layout[0][1][1] <= 195 and layout[0][1][3] >= 345


def ocr_from_pdf(raw, replace=None):
    """OCR stand-in: the PDF's own words at their rendered positions (72 dpi => scale 1)."""
    with fitz.open(stream=raw, filetype="pdf") as pdf:
        words = pdf[0].get_text("words")
    replace = replace or {}
    return {
        "engine": "tesseract",
        "version": "test",
        "status": "complete",
        "words": [
            {
                "text": replace.get(w[4], w[4]),
                "bbox_px": [round(v) for v in w[:4]],
                "confidence": 96.0,
                "line": [w[5], 0, w[6]],
                "order": i,
            }
            for i, w in enumerate(words)
        ],
    }


def vision_reads(text, *, fail=False):
    async def read(image, settings, prompt):
        if fail:
            return {"raw": "The image shows a number.", "finish_reason": "stop"}
        return {
            "raw": json.dumps({"kind": "paragraph", "text": text, "complete": True}),
            "finish_reason": "stop",
        }

    return read


@pytest.fixture
def ocr_settings(settings):
    return settings.model_copy(update={"extraction_ocr_enabled": True})


async def test_text_layer_and_ocr_outvote_a_vision_misread(reader, ocr_settings, monkeypatch):
    raw = pdf_bytes("Value -12.50")

    async def ocr(image, settings):
        return ocr_from_pdf(raw)

    monkeypatch.setattr(hybrid, "ocr_page", ocr)
    reader.side_effect = vision_reads("Value 12.50")  # vision drops the minus sign
    package = await hybrid.extract_document(raw, "source.pdf", ocr_settings)
    block = package.record["pages"][0]["blocks"][0]
    assert block["consensus"] == "native_ocr"
    assert block["selected_text"] == "Value -12.50"
    assert block["proposed_structure"]["text"] == "Value -12.50"
    assert block["ocr_text"] == "Value -12.50" and block["audit_notes"]
    assert "NUMERIC_CONFLICT" not in {i["code"] for i in package.issues}


async def test_text_layer_and_ocr_settle_a_failed_vision_reading(reader, ocr_settings, monkeypatch):
    raw = pdf_bytes("Value -12.50")

    async def ocr(image, settings):
        return ocr_from_pdf(raw)

    monkeypatch.setattr(hybrid, "ocr_page", ocr)
    reader.side_effect = vision_reads("", fail=True)
    package = await hybrid.extract_document(raw, "source.pdf", ocr_settings)
    leaves = [b for b in package.record["pages"][0]["blocks"] if b["status"] != "split"]
    settled = [b for b in leaves if b.get("consensus") == "native_ocr" and b.get("native_text")]
    assert settled and all(b["selected_text"] for b in settled)
    settled_ids = {b["id"] for b in settled}
    assert not any(
        i["code"] == "VISION_UNAVAILABLE" and i.get("block_id") in settled_ids for i in package.issues
    )


async def test_three_way_disagreement_stays_a_review_issue(reader, ocr_settings, monkeypatch):
    raw = pdf_bytes("Value -12.50")

    async def ocr(image, settings):
        return ocr_from_pdf(raw, replace={"-12.50": "-12.80"})

    monkeypatch.setattr(hybrid, "ocr_page", ocr)
    reader.side_effect = vision_reads("Value 12.50")
    package = await hybrid.extract_document(raw, "source.pdf", ocr_settings)
    block = package.record["pages"][0]["blocks"][0]
    assert block.get("selected_text") is None and "consensus" not in block
    assert block["proposed_structure"]["text"] == "Value -12.50"  # draft is the text layer, still undecided
    conflict = next(i for i in package.issues if i["code"] == "NUMERIC_CONFLICT")
    assert conflict["message"] == "PDF text, OCR and vision readings differ"


async def test_unavailable_ocr_is_reported_without_blocking(reader, ocr_settings, monkeypatch):
    async def ocr(image, settings):
        return {
            "engine": "tesseract",
            "status": "unavailable",
            "error": "TesseractNotFoundError",
            "words": [],
        }

    monkeypatch.setattr(hybrid, "ocr_page", ocr)
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", ocr_settings)
    issue = next(i for i in package.issues if i["code"] == "OCR_UNAVAILABLE")
    assert not issue["technical"]
    assert package.record["pages"][0]["blocks"][0]["consensus"] == "native_vision"


async def test_figures_get_a_separate_model_description(reader, settings):
    async def read(image, settings, prompt):
        if prompt == regions.DESCRIPTION_PROMPT:
            return {
                "raw": "A tank feeds a filter, then a pump, which supplies three zones.",
                "finish_reason": "stop",
            }
        return {
            "raw": json.dumps(
                {"kind": "figure", "text": "Tank T-01", "complete": True, "figure_labels": ["Tank T-01"]}
            ),
            "finish_reason": "stop",
        }

    reader.side_effect = read
    page = await regions.read_page(png_bytes(), 1, settings, layout=[("figure", [0, 0, 400, 300])])
    block = page["blocks"][0]
    assert block["figure_description"].startswith("A tank feeds a filter")
    assert block["figure_description_source"] == "vision_model_generated"
    assert block["visual"]["text"] == "Tank T-01"  # description never replaces source text


def test_equivalent_text_ignores_typography_but_not_values():
    from northstar_api.services.extraction_validation import equivalent_text

    assert equivalent_text("“Rate” − 12.5 – kg", '"Rate" - 12.5 - kg')
    assert not equivalent_text("12.50", "12.80")
    assert not equivalent_text("Value -12.50", "Value 12.50")


def test_ocr_lines_keep_source_line_breaks():
    from northstar_api.services.extraction_ocr import lines_text

    words = [
        {"text": t, "line": line, "order": i}
        for i, (t, line) in enumerate(
            [("Zone", [1, 1, 1]), ("Crop", [1, 1, 1]), ("A1", [1, 1, 2]), ("Tomato", [1, 1, 2])]
        )
    ]
    assert lines_text(words) == "Zone Crop\nA1 Tomato"


def test_ocr_confirms_text_layer_only_across_lookalike_glyphs():
    from northstar_api.services.extraction_validation import ocr_confirms

    assert ocr_confirms("A1 132 7 125 REVIEW", "Al 132 7 125 REVIEW")
    assert ocr_confirms("2,240 L/day × 7 days", "2,240 L/day x 7 days")
    assert ocr_confirms("SEP 2026 2 / 2", "SEP 2026 2/2")
    assert not ocr_confirms("Total 285 15 270 -", "Total 285 15 270 2")
    assert not ocr_confirms("53 3 50", "53 8 50")
    assert not ocr_confirms("10 5", "105")
    assert not ocr_confirms("", "")


async def test_region_ocr_pass_confirms_text_the_page_pass_missed(reader, ocr_settings, monkeypatch):
    raw = pdf_bytes("Value -12.50")

    async def page_ocr(image, settings):
        return ocr_from_pdf(raw, replace={"Value": "Va1ue?"})  # page-level OCR garbles a word

    async def region_ocr(image, box, settings):
        return {"text": "Value -12.50", "words": []}

    monkeypatch.setattr(hybrid, "ocr_page", page_ocr)
    monkeypatch.setattr(reconcile, "ocr_region", region_ocr)
    reader.side_effect = vision_reads("Value 12.50")
    package = await hybrid.extract_document(raw, "source.pdf", ocr_settings)
    block = package.record["pages"][0]["blocks"][0]
    assert block["consensus"] == "native_ocr" and block["ocr_region_text"] == "Value -12.50"
    assert block["audit_notes"][0].startswith("A region-level OCR pass confirms the PDF text layer")


async def test_failed_vision_without_ocr_agreement_is_a_decision_not_a_blocker(
    reader, ocr_settings, monkeypatch
):
    raw = pdf_bytes("Total 285 15 270 -")

    async def page_ocr(image, settings):
        return ocr_from_pdf(raw, replace={"-": "2"})

    async def region_ocr(image, box, settings):
        return None

    monkeypatch.setattr(hybrid, "ocr_page", page_ocr)
    monkeypatch.setattr(reconcile, "ocr_region", region_ocr)
    reader.side_effect = vision_reads("", fail=True)
    package = await hybrid.extract_document(raw, "source.pdf", ocr_settings)
    leaves = [b for b in package.record["pages"][0]["blocks"] if b["status"] != "split"]
    drafted = [b for b in leaves if b.get("proposed_structure") and not b.get("consensus")]
    assert drafted
    ids = {b["id"] for b in drafted}
    region_issues = [i for i in package.issues if i.get("block_id") in ids]
    assert region_issues and not any(i["technical"] for i in region_issues)


async def test_failed_vision_on_an_image_with_only_ocr_text_is_a_decision(reader, ocr_settings, monkeypatch):
    async def page_ocr(image, settings):
        return {
            "engine": "tesseract",
            "version": "test",
            "status": "complete",
            "words": [
                {
                    "text": "Week",
                    "bbox_px": [30, 30, 90, 50],
                    "confidence": 90.0,
                    "line": [1, 1, 1],
                    "order": 0,
                }
            ],
        }

    monkeypatch.setattr(hybrid, "ocr_page", page_ocr)

    async def read(image, settings, prompt):
        if prompt == regions.INVENTORY_PROMPT:
            return {
                "raw": json.dumps(
                    {"regions": [{"kind": "figure", "bbox": [0, 0, 1000, 1000]}], "complete": True}
                ),
                "finish_reason": "stop",
            }
        return {"raw": "not json", "finish_reason": "stop"}

    reader.side_effect = read
    package = await hybrid.extract_document(png_bytes(), "chart.png", ocr_settings)
    leaves = [b for b in package.record["pages"][0]["blocks"] if b["status"] != "split"]
    drafted = [b for b in leaves if (b.get("proposed_structure") or {}).get("text") == "Week"]
    assert drafted
    single = [i for i in package.issues if i["code"] == "SINGLE_READER" and i["block_id"] == drafted[0]["id"]]
    assert single and not single[0]["technical"]
    assert not any(
        i["code"] == "VISION_UNAVAILABLE" and i["block_id"] == drafted[0]["id"] for i in package.issues
    )


async def test_vision_is_skipped_where_text_layer_and_ocr_already_agree(reader, ocr_settings, monkeypatch):
    raw = pdf_bytes("Value -12.50")

    async def ocr(image, settings):
        return ocr_from_pdf(raw)

    monkeypatch.setattr(hybrid, "ocr_page", ocr)
    reader.side_effect = vision_reads("", fail=True)  # would fail and split the region if it were called
    package = await hybrid.extract_document(raw, "source.pdf", ocr_settings)
    blocks = package.record["pages"][0]["blocks"]
    assert reader.await_count == 0  # 2 of 3 already met: no vision call, no split, no boundary check
    assert not any(b["status"] == "split" for b in blocks)
    assert all(b["consensus"] == "native_ocr" for b in blocks)
    codes = {i["code"] for i in package.issues}
    assert "SPLIT_BOUNDARY_REVIEW" not in codes and "NUMERIC_CONFLICT" not in codes


async def test_split_image_regions_still_get_a_boundary_check(reader, settings):
    calls = {"n": 0}

    async def read(image, settings, prompt):
        calls["n"] += 1
        if prompt == regions.INVENTORY_PROMPT:
            return {
                "raw": json.dumps(
                    {"regions": [{"kind": "paragraph", "bbox": [0, 0, 1000, 1000]}], "complete": True}
                ),
                "finish_reason": "stop",
            }
        if calls["n"] == 2:  # the full-page crop fails once, its halves are read
            return {"raw": "not json", "finish_reason": "stop"}
        return {
            "raw": json.dumps({"kind": "paragraph", "text": "half", "complete": True}),
            "finish_reason": "stop",
        }

    reader.side_effect = read
    package = await hybrid.extract_document(png_bytes(), "scan.png", settings)
    assert "SPLIT_BOUNDARY_REVIEW" in {i["code"] for i in package.issues}
